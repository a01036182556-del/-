"""
Shared helpers for the 건설기계 기초자료 관리대장 (construction equipment ledger).

Two entry points:
  - create_ledger(output_path): build a brand-new blank ledger workbook.
  - append_row(xlsx_path, values, comments=None): add one equipment's data
    as a new row in an existing ledger, auto-numbering 연번 and auto-growing
    the sheet if the pre-built blank rows have all been used.

`values` and `comments` are dicts keyed by column NUMBER (1-22, see COLUMNS
below) rather than column letter, so the calling script stays readable even
though several headers are long Korean phrases.

Column numbers (also printed by `print_columns()`):
  1 연번            8  소유자(임대업체)   15 보험기간
  2 장비명(기종)     9  대표자             16 운전원 성명
  3 규격/모델명      10 연락처             17 면허종류
  4 등록번호         11 검사종류           18 면허번호
  5 차대(형식)번호    12 검사유효기간      19 면허유효기간
  6 제작사           13 보험회사          20 배치현장
  7 제작연도         14 보험증권번호      21 반입일
                                          22 비고
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from datetime import date

FONT_NAME = "맑은 고딕"
SHEET_TITLE = "건설기계 기초자료 관리대장"

# (header, group, width) -- order defines the column number (1-indexed)
COLUMNS = [
    ("연번", "기본정보", 5),
    ("장비명(기종)", "기본정보", 14),
    ("규격/모델명", "기본정보", 14),
    ("등록번호", "기본정보", 14),
    ("차대(형식)번호", "기본정보", 16),
    ("제작사", "기본정보", 12),
    ("제작연도", "기본정보", 9),
    ("소유자(임대업체)", "소유/연락처", 16),
    ("대표자", "소유/연락처", 10),
    ("연락처", "소유/연락처", 14),
    ("검사종류", "검사현황", 10),
    ("검사유효기간", "검사현황", 14),
    ("보험회사", "보험현황", 12),
    ("보험증권번호", "보험현황", 16),
    ("보험기간", "보험현황", 20),
    ("운전원 성명", "운전원정보", 10),
    ("면허종류", "운전원정보", 12),
    ("면허번호", "운전원정보", 16),
    ("면허유효기간", "운전원정보", 14),
    ("배치현장", "현장정보", 14),
    ("반입일", "현장정보", 11),
    ("비고", "현장정보", 16),
]
N_COLS = len(COLUMNS)

# Columns that read better left-aligned (free text) vs. centered (codes/dates)
LEFT_COLS = {2, 3, 8, 20, 22}

_HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_GROUP_FILL = PatternFill("solid", fgColor="8EA9DB")
_EXAMPLE_FONT = Font(name=FONT_NAME, size=10, color="0000FF")
_NORMAL_FONT = Font(name=FONT_NAME, size=10)
_LEGEND_FONT = Font(name=FONT_NAME, size=9, color="808080", italic=True)
_TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True)
_THIN = Side(style="thin", color="B7B7B7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

TITLE_ROW = 1
NOTE_ROW = 2
GROUP_ROW = 3
HEADER_ROW = 4
EXAMPLE_ROW = 5
FIRST_BLANK_ROW = 6
N_BLANK_ROWS = 25  # pre-built empty rows a fresh ledger ships with


def _cell_alignment(col):
    return Alignment(
        horizontal="left" if col in LEFT_COLS else "center",
        vertical="center",
        wrap_text=True,
    )


def print_columns():
    for i, (header, group, _w) in enumerate(COLUMNS, start=1):
        print(f"{i:2d}  [{group}] {header}")


def create_ledger(output_path):
    """Build a brand-new blank ledger workbook and save it to output_path."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=N_COLS)
    ws.cell(row=TITLE_ROW, column=1, value=SHEET_TITLE).font = _TITLE_FONT
    ws.cell(row=TITLE_ROW, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[TITLE_ROW].height = 28

    ws.merge_cells(start_row=NOTE_ROW, start_column=1, end_row=NOTE_ROW, end_column=N_COLS)
    ws.cell(
        row=NOTE_ROW,
        column=1,
        value=f"작성일: {date.today().isoformat()}   "
        f"※ 파란색 글자 행({EXAMPLE_ROW}행)은 작성 예시이므로 실제 장비 등록 시 삭제 후 사용하세요.",
    ).font = _LEGEND_FONT
    ws.cell(row=NOTE_ROW, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[NOTE_ROW].height = 16

    # Group header row: merge consecutive columns sharing the same group label
    c = 1
    while c <= N_COLS:
        grp = COLUMNS[c - 1][1]
        start = c
        while c <= N_COLS and COLUMNS[c - 1][1] == grp:
            c += 1
        end = c - 1
        if end > start:
            ws.merge_cells(start_row=GROUP_ROW, start_column=start, end_row=GROUP_ROW, end_column=end)
        for cc in range(start, end + 1):
            cell = ws.cell(row=GROUP_ROW, column=cc, value=grp if cc == start else None)
            cell.font = _HEADER_FONT
            cell.fill = _GROUP_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER
    ws.row_dimensions[GROUP_ROW].height = 18

    for idx, (label, _grp, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[HEADER_ROW].height = 30

    example_values = [
        1, "굴착기", "0.7m3(SOLAR 220N-V)", "경기12가1234", "HD1234567",
        "현대건설기계", 2019, "(주)대한중기임대", "홍길동", "010-1234-5678",
        "정기검사", "2026-11-30", "DB손해보험", "1234567890",
        "2026-01-01 ~ 2026-12-31", "김철수", "굴착기운전기능사", "12-345678-90",
        "2027-05-20", "○○현장 신축공사", "2026-08-25", "예시 행 - 삭제 후 사용",
    ]
    for idx, val in enumerate(example_values, start=1):
        cell = ws.cell(row=EXAMPLE_ROW, column=idx, value=val)
        cell.font = _EXAMPLE_FONT
        cell.alignment = _cell_alignment(idx)
        cell.border = _BORDER
    ws.row_dimensions[EXAMPLE_ROW].height = 18

    last_blank = FIRST_BLANK_ROW + N_BLANK_ROWS - 1
    for r in range(FIRST_BLANK_ROW, last_blank + 1):
        for idx in range(1, N_COLS + 1):
            cell = ws.cell(row=r, column=idx)
            cell.font = _NORMAL_FONT
            cell.border = _BORDER
            cell.alignment = _cell_alignment(idx)
        ws.cell(row=r, column=1, value=r - FIRST_BLANK_ROW + 1)
        ws.row_dimensions[r].height = 16

    ws.freeze_panes = f"B{HEADER_ROW + 1}"
    ws.sheet_view.showGridLines = False

    wb.save(output_path)
    return output_path


def _find_header_row(ws):
    """Locate the header row by scanning column A for the '연번' label,
    so this keeps working even if a ledger was hand-edited a bit."""
    for r in range(1, 10):
        if ws.cell(row=r, column=1).value == "연번":
            return r
    raise ValueError("'연번' 헤더를 찾지 못했습니다 — 이 파일이 이 스킬로 만든 관리대장이 맞는지 확인하세요.")


def append_row(xlsx_path, values, comments=None, save=True):
    """Append one equipment's data to an existing ledger.

    values:   {column_number: value, ...} — only include columns you actually
              have data for; anything omitted is left blank. Column 1 (연번)
              is filled in automatically — don't pass it.
    comments: {column_number: comment_text, ...} — use this for any field
              you filled with a hedge like "판독 어려움" or "-" so the reason
              is visible to whoever opens the sheet next (see SKILL.md's
              "never fabricate illegible data" rule).

    Returns the row number that was written to.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    header_row = _find_header_row(ws)
    example_row = header_row + 1
    first_blank = header_row + 2

    # Find the first row (from first_blank onward) whose 장비명(기종) cell
    # (column 2) is still empty -- column 1 is pre-numbered even when blank,
    # so it can't be used as the "is this row free" check.
    row = first_blank
    while ws.cell(row=row, column=2).value not in (None, ""):
        row += 1

    # Ran out of pre-built blank rows -- extend the sheet with one more,
    # copying the style of the last existing row.
    max_row = ws.max_row
    if row > max_row:
        for col in range(1, N_COLS + 1):
            src = ws.cell(row=max_row, column=col)
            dst = ws.cell(row=row, column=col)
            dst.font = src.font
            dst.border = src.border
            dst.alignment = src.alignment
        ws.row_dimensions[row].height = ws.row_dimensions[max_row].height

    seq = row - first_blank + 1
    ws.cell(row=row, column=1, value=seq)

    for col, val in values.items():
        if col == 1:
            continue  # 연번 is auto-assigned above
        cell = ws.cell(row=row, column=col, value=val)
        cell.alignment = _cell_alignment(col)

    if comments:
        for col, text in comments.items():
            ws.cell(row=row, column=col).comment = Comment(text, "Claude")

    if save:
        wb.save(xlsx_path)
    return row


if __name__ == "__main__":
    import sys

    if len(sys.argv) == 3 and sys.argv[1] == "create":
        create_ledger(sys.argv[2])
        print(f"created {sys.argv[2]}")
    else:
        print(__doc__)
        print_columns()
